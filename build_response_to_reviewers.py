from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(r"E:\project\xiaozhi\Uni-Sign")
TEMPLATE = ROOT / "docs" / "04_论文稿件" / "回复编辑论文内容" / "_ IEEE-Access-Response-to-Reviewers-template-6.26.23.docx"
OUT = ROOT / "docs" / "04_论文稿件" / "回复编辑论文内容" / "Response_to_Reviewers_Access-2024-38112.docx"

doc = Document(TEMPLATE)
body = doc._element.body
for child in list(body):
    if child.tag != qn("w:sectPr"):
        body.remove(child)

section = doc.sections[0]
section.top_margin = Inches(0.75)
section.bottom_margin = Inches(0.75)
section.left_margin = Inches(0.82)
section.right_margin = Inches(0.82)

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Times New Roman"
normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
normal.font.size = Pt(10.5)
normal.paragraph_format.space_after = Pt(5)
normal.paragraph_format.line_spacing = 1.08

if "Title" not in [s.name for s in styles]:
    title_style = styles.add_style("Title", WD_STYLE_TYPE.PARAGRAPH)
else:
    title_style = styles["Title"]
title_style.base_style = normal
title_style.font.name = "Times New Roman"
title_style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
title_style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
title_style.font.size = Pt(15)
title_style.font.bold = True
title_style.font.color.rgb = RGBColor(0, 0, 0)
title_style.paragraph_format.space_after = Pt(8)

for name, size, bold, before, after in [("Response Heading", 13, True, 12, 5), ("Reviewer Heading", 12, True, 14, 6)]:
    if name not in [s.name for s in styles]:
        style = styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
    else:
        style = styles[name]
    style.base_style = normal
    style.font.name = "Times New Roman"
    style._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    style._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.paragraph_format.space_before = Pt(before)
    style.paragraph_format.space_after = Pt(after)
    style.paragraph_format.keep_with_next = True

def add(text="", style=None, align=None, italic=False, bold_lead=None):
    p = doc.add_paragraph(style=style)
    if align is not None:
        p.alignment = align
    if bold_lead and text.startswith(bold_lead):
        r = p.add_run(bold_lead)
        r.bold = True
        r = p.add_run(text[len(bold_lead):])
        r.italic = italic
    else:
        r = p.add_run(text)
        r.italic = italic
    return p

def label_paragraph(label, text, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.15)
    p.paragraph_format.first_line_indent = Inches(-0.15)
    p.paragraph_format.space_after = Pt(5)
    r = p.add_run(label)
    r.bold = True
    r = p.add_run(text)
    r.italic = italic
    return p

def concern(number, quote, response, action):
    add(f"Concern {number}", style="Response Heading")
    label_paragraph("Reviewer's concern: ", f'“{quote}”', italic=True)
    label_paragraph("Author response: ", response)
    label_paragraph("Author action: ", action)

add("Response to Reviewers", style="Title", align=WD_ALIGN_PARAGRAPH.CENTER)
add("Manuscript ID: Access-2024-38112", align=WD_ALIGN_PARAGRAPH.CENTER)
add("Manuscript title: When Does RGB Help? A Fine-Grained Empirical Analysis of RGB Fusion for Chinese Sign Language Translation on CE-CSL", align=WD_ALIGN_PARAGRAPH.CENTER)

add("To: IEEE Access Editor", bold_lead="To: ")
add("Re: Response to reviewers", bold_lead="Re: ")
add()
add("Dear Editor,")
add("We sincerely thank you and the reviewers for the constructive and detailed comments on our manuscript. We have carefully revised the manuscript and conducted additional analyses to address every concern. The revision clarifies the scope of our diagnostic study, strengthens the experimental evidence, and improves the reproducibility and presentation of the work. The principal additions are a ten-seed evaluation with paired statistical tests, a gate-weight distribution analysis, complete architecture and training details, clarified metric definitions and table values, and a transparent comparison framework for CE-CSL.")
add("For clarity, each response below contains the reviewer’s concern, our response, and the corresponding action in the revised manuscript. The concern wording is reproduced from the review summary available to us; it should be checked against the original reviewer text in the submission system before upload.")
add("We respectfully submit the revised manuscript and this point-by-point response for your consideration.")
add("Sincerely,")
add("Zhaoyong Fan et al.")

add("Reviewer 1", style="Reviewer Heading")
concern(1,
        "Please provide a multi-seed evaluation and report the statistical reliability of the conclusions.",
        "We agree that a single run is insufficient in this low-resource setting. We therefore expanded the original three-seed study (42, 123, and 3407) to ten random seeds by adding 456, 789, 101, 202, 303, 404, and 505. The pose-only model achieved the highest AvgSim in every paired comparison with both RGB variants. Across the ten runs, pose-only obtained 0.3443 ± 0.0121, compared with 0.3309 ± 0.0160 for naive RGB and 0.3304 ± 0.0152 for the full model. A paired t-test gives p = 0.0013 for full model versus pose-only; the Wilcoxon signed-rank test gives the same ranking, and a 10,000-resample bootstrap confidence interval for the paired difference excludes zero. We revised the claims accordingly: RGB does not reliably improve the aggregate result under the tested configuration.",
        "Added the expanded ten-seed protocol, per-seed results, standard deviations, paired tests, Wilcoxon check, and bootstrap interval in Section V-D (Multi-Seed Stability Analysis) and Table 6; updated the Abstract, Introduction, Discussion, and Conclusion to reflect the multi-seed evidence.")
concern(2,
        "Please analyze the influence of the RGB encoder/backbone and clarify whether the conclusions are backbone-dependent.",
        "We appreciate this important request. The revised manuscript now identifies EfficientNet-B0, initialized from ImageNet weights, as the RGB encoder and reports its 1280-dimensional output and approximately 5.3M parameters. A 1 × 1 convolution projects the RGB representation from 1280 to 256 dimensions. We explicitly state that the pose encoder and mT5 decoder are frozen in all ablations, so the comparison isolates the RGB-support path. We also clarify that the conclusions are limited to CE-CSL with MediaPipe v0.10.32 and EfficientNet-B0; they do not establish a result for other RGB backbones, languages, or architectures. Replacing EfficientNet-B0 with a smaller encoder is identified as future work, not claimed as an evaluated result.",
        "Added component-level architecture and parameter counts in Table 1; clarified the RGB encoder, projection, frozen components, and scope of inference in Sections III-D, III-E, IV-D, V-H, and V-L.")
concern(3,
        "Please provide the distribution of gate weights and explain what the learned fusion gate actually does.",
        "We added a gate analysis that reports the learned RGB weight α_t rather than relying only on aggregate scores. The results show that the gate remains conservative and close to its pose-biased initialization for much of the sequence. This observation is interpreted cautiously: it is consistent with limited data, noisy RGB crops, and a capacity-data mismatch, but it does not establish a causal mechanism. The revised discussion distinguishes this descriptive evidence from the statistically supported ten-seed comparison.",
        "Added the gate-weight distribution/behavior analysis in Section V-K (Interpretive Analysis and Discussion), with the gate definition and initialization retained in Sections III-D and III-E; linked the interpretation to the ablation results in Table 8.")
concern(4,
        "Please state the sample size and training details sufficiently for reproducibility.",
        "We have expanded the experimental specification. CE-CSL contains 5,988 clips, with 4,973 training samples, 515 development samples, and 500 test samples; Dev100 and Test100 are the 100-example evaluation subsets used for the principal comparisons. The revision also reports the number of backgrounds, vocabulary-related statistics, sentence-length statistics, the four training ratios, optimizer, learning-rate schedule, warmup, weight decay, label smoothing, effective batch size, maximum sequence length, precision, initialization, and approximate run time. The trainable additions to the RGB path total approximately 6.0M parameters, while the shared pose encoder and mT5 decoder remain frozen.",
        "Expanded Sections IV-A, IV-D, and IV-E; added dataset statistics in Table 2 and component-level parameter accounting in Table 1 and Table 10.")

add("Reviewer 2", style="Reviewer Heading")
concern(1,
        "Please clarify the role of the mT5 decoder in Uni-Sign and in the oracle gloss-to-text comparison.",
        "Thank you for noting that the decoder role needed clearer explanation. The revised manuscript states that Uni-Sign uses an mT5 decoder for multilingual text generation. In every RGB-pose ablation, the same pretrained mT5 decoder is frozen; only the RGB-side support components differ. The oracle gloss-to-text reference uses the same mT5 decoder configuration but receives ground-truth gloss annotations rather than continuous visual input. Accordingly, it is presented as an approximate upper bound for the gloss-to-text path, not as a directly comparable end-to-end visual model. The observed oracle gap is discussed as evidence that visual-to-semantic encoding is a major source of error in this configuration, without attributing the gap to a single cause.",
        "Clarified the Uni-Sign/mT5 architecture in Sections II-C and III-A; specified the oracle setup and its interpretation in Section IV-C, Table 4, and Section V-B.")
concern(2,
        "Please verify the BLEU-4 and WER metrics and resolve the apparent inconsistency in Table 3.",
        "We carefully rechecked the metric definitions, evaluation outputs, and all values reported in Table 3. BLEU-4 (B-4), ROUGE-L, chrF, and WER are reported together with AvgSim; higher is better for AvgSim, BLEU-4, ROUGE-L, and chrF, whereas lower is better for WER. The revised text makes clear that the main six-variant comparison and the separate ten-seed experiment answer different questions and should not be conflated. In Table 3, the Test100 full model records AvgSim 0.3620, BLEU-4 10.62, and WER 74.91; pose-only records 0.3625, 12.34, and 73.70, respectively. The metadata-free gate records the best BLEU-4 among RGB variants (11.72) but does not exceed pose-only. We now explicitly discuss this metric-dependent pattern rather than implying a uniform winner.",
        "Verified and clarified metric directions and Table 3 values in Sections IV-B and V-A; added the metric-dependent interpretation in Sections V-A and V-F, with the focused RGB ablation presented in Table 8.")
concern(3,
        "Please justify the Top-K 10% selection rule and state its implementation details and limitations.",
        "We have clarified that the Top-K ratio was a heuristic design choice, not an optimized hyperparameter or a claim of universal optimality. After discarding crops with hand confidence below 0.3, we retain K = ceil(0.1N) candidate crops, while keeping at least one frame whenever candidates are available. The selected RGB features are aligned to the pose timeline through the stated alignment procedure. We also explicitly acknowledge that a controlled sensitivity sweep over 5%, 10%, 15%, and 20% is needed before making a stronger claim about the ratio; this experiment is presented as future work.",
        "Expanded the selection rule and edge-case handling in Sections III-B and IV-D; added the limitation and planned sensitivity analysis in Section V-L.")

add("Reviewer 3", style="Reviewer Heading")
concern(1,
        "Please explain the meaning and scope of the title ‘When Does RGB Help?’",
        "We agree that the title requires a precise, evidence-based interpretation. The revised manuscript explains that the question is conditional rather than a claim of a universal RGB benefit. In the aggregate ten-seed analysis, pose-only is more accurate and more stable than both evaluated RGB variants. The title refers to the narrower exploratory sentence-level pattern: on Test100, the full model achieved the highest AvgSim for long sentences (0.3619, +0.0105 relative to pose-only) and for questions. These subgroup observations are explicitly labelled exploratory because Test100 contains 100 samples and the category counts are limited. The title is therefore aligned with the manuscript’s central diagnostic question, while the Abstract, Introduction, Discussion, and Conclusion state the limits of the evidence.",
        "Reframed the title’s scope and the conditional conclusion in the Abstract, Introduction, Section V-G, Section V-K, and Conclusion; reported the subgroup counts and values in Table 9.")
concern(2,
        "Please add an appropriate comparison with state-of-the-art methods on the CE-CSL dataset.",
        "We revised the comparison to avoid an invalid cross-dataset leaderboard. Table 5 now separates two comparison types: published SLT/SLR methods evaluated on different datasets and metrics, and the directly comparable CE-CSL Test100 results obtained with the same Uni-Sign backbone and split. The latter includes the oracle gloss-to-text reference, pose-only, the metadata-free RGB gate, and the full model. Because published methods do not provide results on the identical CE-CSL split under the same protocol, we do not claim state-of-the-art performance. Instead, the manuscript states clearly that this is a controlled diagnostic study and that cross-paper numbers serve as a design-context map only.",
        "Reorganized the external comparison in Section V-C and Table 5; strengthened the non-SOTA, controlled-study framing in the Abstract, Introduction, and Conclusion.")
concern(3,
        "Please correct the abbreviation and formatting in Table 5.",
        "We have standardized the abbreviations and presentation in Table 5. The table now clearly distinguishes the method, venue, data, input, and design focus, and it labels the CE-CSL entries as directly comparable only when they share the same backbone and Test100 split. The caption also states that the upper portion contains different datasets and metrics, preventing readers from interpreting it as a common leaderboard.",
        "Corrected Table 5 abbreviations, labels, and explanatory caption; revised the accompanying text in Section V-C for consistency.")

add("Closing statement", style="Response Heading")
add("We are grateful for the reviewers’ guidance. We believe that the revised manuscript now gives a more rigorous, transparent, and reproducible account of the conditions under which RGB support was evaluated on CE-CSL. All conclusions have been narrowed to the evidence provided, and no claim of general backbone-, dataset-, or language-independent improvement is made.")

doc.core_properties.title = "Response to Reviewers Access 2024 38112"
doc.core_properties.subject = "Point-by-point response to reviewers"
doc.core_properties.author = ""
doc.save(OUT)
print(OUT)
