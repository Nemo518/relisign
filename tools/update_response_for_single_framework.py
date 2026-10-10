from pathlib import Path
import sys

from docx import Document


source = Path(sys.argv[1])
output = Path(sys.argv[2])

replacements = {
    "Author action: We removed the disputed qualitative figure. The revised figures are method diagrams only and do not claim to visualize a particular test sample.":
        "Author action: We removed the disputed qualitative figure. The revised Figure 1 is a method diagram only and does not claim to visualize a particular test sample.",
    "Manuscript location: Pages 5--6; Figures 1--2 are method diagrams. The former qualitative sample figure and dependent discussion were removed.":
        "Manuscript location: Pages 4--6; Figure 1 is the method diagram. The former qualitative sample figure and dependent discussion were removed.",
    "Manuscript location: Page 10; Section V-F (Gate Distribution).":
        "Manuscript location: Page 9; Section V-F (Gate Distribution).",
}

document = Document(source)
found = set()
for paragraph in document.paragraphs:
    if paragraph.text not in replacements:
        continue
    replacement = replacements[paragraph.text]
    if not paragraph.runs:
        paragraph.add_run(replacement)
    elif len(paragraph.runs) == 1:
        paragraph.runs[0].text = replacement
    else:
        prefix = paragraph.runs[0].text
        if replacement.startswith(prefix):
            paragraph.runs[1].text = replacement[len(prefix):]
            for run in paragraph.runs[2:]:
                run.text = ""
        else:
            paragraph.runs[0].text = replacement
            for run in paragraph.runs[1:]:
                run.text = ""
    found.add(paragraph.text)

if len(found) != len(replacements):
    missing = [key for key in replacements if not any(p.text == replacements[key] for p in document.paragraphs)]
    if missing:
        raise RuntimeError(f"Missing replacement targets: {missing}")

output.parent.mkdir(parents=True, exist_ok=True)
document.save(output)
print(output)
